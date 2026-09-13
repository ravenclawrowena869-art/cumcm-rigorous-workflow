"""Read-only Phase A attachment adapters; no forecast/actual substitution.

Actual ``actual_available_at`` is a REPLAY_SLOT_ENDPOINT convention: the
just-ended ten-minute slot is exposed before the boundary's stage decision.
The source does not contain observed telemetry availability timestamps.
Forecast ``known_at = issue_time`` represents the attachment's stated release
schedule, not independently observed ingestion metadata. All times use +08:00.
Attachment 1 contributes only its fixed tariff, never its load/PV columns.
"""
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import math
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TZ = timezone(timedelta(hours=8))
ACTUAL_AVAILABILITY_CONVENTION = "REPLAY_SLOT_ENDPOINT_NOT_OBSERVED_TELEMETRY"
FORECAST_AVAILABILITY_CONVENTION = "STATED_ISSUE_TIME_NOT_OBSERVED_INGESTION"


def _invalid(where, reason):
    raise ValueError(f"INPUT_INVALID: {where}: {reason}")


def _read(path):
    """Hash and parse the same file handle; close workbook even on read errors."""
    path = Path(path)
    with path.open("rb") as source:
        source_hash = hashlib.sha256(source.read()).hexdigest()
        source.seek(0)
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheets = {sheet.title: list(sheet.values) for sheet in workbook}
            finally:
                workbook.close()
        except (BadZipFile, ValueError, TypeError, KeyError) as exc:
            _invalid(path.name, f"unreadable workbook ({exc})")
    return sheets, source_hash


def _shape(rows, count, columns, where):
    if len(rows) != count + 1 or any(len(row) != columns for row in rows):
        _invalid(where, f"expected one header and {count} rows / {columns} columns")


def _number(value, where):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        _invalid(where, "missing or nonnumeric value")
    if not math.isfinite(value) or value < 0:
        _invalid(where, "power/price must be finite and nonnegative")
    return float(value)


def _minutes(value, where):
    """Interpret explicit day-end labels without folding them into 00:00."""
    if isinstance(value, time):
        if value.second or value.microsecond or value.tzinfo is not None:
            _invalid(where, "invalid time label")
        return 60 * value.hour + value.minute
    if isinstance(value, str):
        if value in ("0:00+1", "00:00+1", "24:00", "24:00:00"):
            return 1440
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?", value)
        if match:
            hour, minute = map(int, match.groups())
            if 0 <= hour <= 23 and 0 <= minute <= 59:
                return hour * 60 + minute
    _invalid(where, "missing or malformed time label")


def _day(value, where):
    if isinstance(value, datetime):
        if value.time() != time() or value.tzinfo is not None:
            _invalid(where, "date must be a naive midnight date")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", value):
        try:
            return date(*map(int, value.split("-")))
        except ValueError:
            pass
    _invalid(where, "missing or malformed date")


def _calendar(days, where):
    if len(days) != 365 or len(set(days)) != 365:
        _invalid(where, "expected 365 distinct dates; missing/duplicate date")
    start = date(days[0].year, 1, 1)
    if days != [start + timedelta(days=i) for i in range(365)] or days[-1] != date(start.year, 12, 31):
        _invalid(where, "expected one complete chronological 365-day calendar year")


def load_tariff(path) -> list[dict]:
    """Return the 144 delivery-slot prices from attachment 1's fixed day."""
    sheets, _ = _read(path)
    if len(sheets) != 1:
        _invalid("tariff", "expected exactly one sheet")
    name, rows = next(iter(sheets.items()))
    _shape(rows, 144, 4, name)
    if rows[0][:2] != ("时间", "电价"):
        _invalid(name, "unexpected time/price headers")
    result = []
    for slot, row in enumerate(rows[1:], 1):
        if _minutes(row[0], f"{name}!A{slot + 1}") != 10 * slot:
            _invalid(f"{name}!A{slot + 1}", "duplicate/missing/out-of-order slot endpoint")
        result.append({
            "slot_id": slot,
            "price_yuan_per_kwh": _number(row[1], f"{name}!B{slot + 1}"),
            "raw_time_label": row[0].isoformat() if isinstance(row[0], time) else row[0],
        })
    return result


def load_actuals(path) -> dict:
    """Load separate actual channels, exposed at endpoints by replay convention.

    These values are replay inputs, never a decision-time load forecast.
    ``ACTUAL_AVAILABILITY_CONVENTION`` explicitly describes the assigned times.
    """
    sheets, _ = _read(path)
    names = ("小区负载", "光伏发电实际功率")
    if set(sheets) != set(names):
        _invalid("actuals", "expected exactly load and actual PV sheets")
    channels = []
    calendars = []
    for name in names:
        rows = sheets[name]
        _shape(rows, 365, 145, name)
        if rows[0][0] != "日期\\时间":
            _invalid(name, "unexpected date header")
        for slot, label in enumerate(rows[0][1:], 1):
            if _minutes(label, f"{name}!{get_column_letter(slot + 1)}1") != slot * 10:
                _invalid(name, "duplicate/missing/out-of-order slot header")
        days = [_day(row[0], f"{name}!A{i}") for i, row in enumerate(rows[1:], 2)]
        _calendar(days, name)
        calendars.append(days)
        channels.append([
            [_number(value, f"{name}!{get_column_letter(j)}{i}")
             for j, value in enumerate(row[1:], 2)]
            for i, row in enumerate(rows[1:], 2)
        ])
    if calendars[0] != calendars[1]:
        _invalid("actuals", "load and PV sheets must have the same dates")
    result = {}
    for day, loads, pvs in zip(calendars[0], channels[0], channels[1]):
        midnight = datetime.combine(day, time(), TZ)
        result[day.isoformat()] = [
            {"slot_id": slot, "load_kw": load, "pv_kw": pv,
             "load_kwh": load / 6, "pv_kwh": pv / 6,
             "actual_available_at": (midnight + timedelta(minutes=10 * slot)).isoformat()}
            for slot, (load, pv) in enumerate(zip(loads, pvs), 1)
        ]
    return result


def load_vintages(path) -> list[dict]:
    """Validate 365 explicit four-stage date blocks, then expose 24 hourly anchors.

    Only the three structural blanks inside each valid block are filled down.
    Targets are issue + 1..24 hours, including legal next-year forecast targets.
    The workbook SHA256 and original one-based Excel row identify the source.
    """
    sheets, source_hash = _read(path)
    if len(sheets) != 1:
        _invalid("vintages", "expected exactly one sheet")
    name, rows = next(iter(sheets.items()))
    _shape(rows, 1460, 26, name)
    expected_header = ("日期", "预报时刻", *(f"预报{i}小时" for i in range(1, 25)))
    if rows[0] != expected_header:
        _invalid(name, "expected date, stage, and ordered 1..24-hour headers")
    days = []
    result = []
    for offset in range(0, 1460, 4):
        block = rows[offset + 1:offset + 5]
        day = _day(block[0][0], f"{name}!A{offset + 2}")
        days.append(day)
        for stage_index, row in enumerate(block):
            source_row = offset + stage_index + 2
            if stage_index and row[0] not in (None, ""):
                _invalid(f"{name}!A{source_row}", "date may appear only at the block's 00 stage")
            stage = stage_index * 6
            if _minutes(row[1], f"{name}!B{source_row}") != stage * 60:
                _invalid(f"{name}!B{source_row}", "block must contain exactly ordered stages 0,6,12,18")
            issue = datetime.combine(day, time(stage), TZ)
            issue_string = issue.isoformat()
            vintage_id = f"{source_hash}:{name}:row{source_row}"
            anchors = [
                {"anchor_id": f"{vintage_id}:h{hour:02d}",
                 "target_time": (issue + timedelta(hours=hour)).isoformat(),
                 "pv_kw": _number(value, f"{name}!{get_column_letter(hour + 2)}{source_row}"),
                 "known_at": issue_string}
                for hour, value in enumerate(row[2:], 1)
            ]
            result.append({"issue_time": issue_string, "known_at": issue_string,
                           "vintage_id": vintage_id, "source_row": source_row,
                           "source_hash": source_hash, "anchors": anchors})
    _calendar(days, name)
    return result
