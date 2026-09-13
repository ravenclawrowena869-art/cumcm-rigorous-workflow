from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

SHEETS = ['计划购电量', '充放电量', '紧急购电量']
PERIODS = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
TOL = 1e-9


def _date_text(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return pd.Timestamp(value).date().isoformat()


def _max_abs(values) -> float:
    return float(max((abs(float(x)) for x in values), default=0.0))


def _emergency_groups(df: pd.DataFrame):
    rows = []
    for day, group in df.groupby('date', sort=True):
        group = group.sort_values('slot')
        active = group[group.emergency_kWh > 1e-9].slot.astype(int).tolist()
        blocks = []
        for slot in active:
            if not blocks or slot != blocks[-1][-1] + 1:
                blocks.append([slot])
            else:
                blocks[-1].append(slot)
        first = True
        for block in blocks:
            start, end = 10 * (block[0] - 1), 10 * block[-1]
            fmt = lambda m: '24:00' if m == 1440 else f'{m // 60}:{m % 60:02d}'
            amount = float(group[group.slot.isin(block)].emergency_kWh.sum())
            rows.append((day if first else None, f'{fmt(start)}-{fmt(end)}', amount))
            first = False
    return rows


def main(slot_csv: Path, workbook: Path, output: Path):
    df = pd.read_csv(slot_csv).sort_values(['date', 'slot']).reset_index(drop=True)
    wb = openpyxl.load_workbook(workbook, data_only=True)
    checks = {
        'causal_source_only': set(df['price_mode'].astype(str)) == {'CAUSAL_LAG7'},
        'sheet_names_exact': wb.sheetnames == SHEETS,
        'slot_coverage': bool(len(df) == 334 * 144 and df.groupby('date').slot.nunique().eq(144).all()),
    }
    if not checks['sheet_names_exact']:
        out = {'readback_verdict': 'FAIL', 'checks': checks, 'actual_sheet_names': wb.sheetnames}
        output.write_text(json.dumps(out, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        raise SystemExit(2)

    purchase, battery, emergency = (wb[name] for name in SHEETS)
    checks['dimensions'] = (
        purchase.max_row == 335 and purchase.max_column == 147 and
        battery.max_row == 1 + 334 * 6 and battery.max_column == 6 and
        emergency.max_column == 3
    )

    q_diffs, q_total_diffs, cost_diffs, date_errors = [], [], [], 0
    for row, (day, group) in enumerate(df.groupby('date', sort=True), 2):
        group = group.sort_values('slot')
        date_errors += int(_date_text(purchase.cell(row, 1).value) != day)
        q_diffs.extend(float(purchase.cell(row, col).value or 0) - float(q)
                       for col, q in enumerate(group.q_DA_kWh, 2))
        q_total_diffs.append(float(purchase.cell(row, 146).value or 0) - float(group.q_DA_kWh.sum()))
        cost_diffs.append(float(purchase.cell(row, 147).value or 0) - float(group.normal_cost_yuan.sum()))

    charge_diffs, discharge_diffs, soc_diffs, battery_label_errors = [], [], [], 0
    row = 2
    for day, group in df.groupby('date', sort=True):
        group = group.sort_values('slot')
        for period, label in enumerate(PERIODS):
            block = group.iloc[period * 24:(period + 1) * 24]
            expected_date = day if period == 0 else None
            actual_date = battery.cell(row, 1).value
            if expected_date is None:
                battery_label_errors += int(actual_date is not None)
            else:
                battery_label_errors += int(_date_text(actual_date) != expected_date)
            battery_label_errors += int(battery.cell(row, 2).value != label)
            charge_diffs.append(float(battery.cell(row, 3).value or 0) - float(block.charge_exec_kWh.sum()))
            discharge_diffs.append(float(battery.cell(row, 4).value or 0) - float(block.discharge_exec_kWh.sum()))
            if period == 0:
                battery_label_errors += int(str(battery.cell(row, 5).value) not in {'0:00', '00:00:00'})
                soc_diffs.append(float(battery.cell(row, 6).value or 0) - float(group.iloc[0].SOC_start_kWh))
            elif period == 1:
                battery_label_errors += int(str(battery.cell(row, 5).value) != '24:00')
                soc_diffs.append(float(battery.cell(row, 6).value or 0) - float(group.iloc[-1].SOC_end_kWh))
            else:
                battery_label_errors += int(battery.cell(row, 5).value is not None or battery.cell(row, 6).value is not None)
            row += 1

    expected_emergency = _emergency_groups(df)
    emergency_amount_diffs, emergency_label_errors = [], 0
    checks['emergency_row_count'] = emergency.max_row == 1 + len(expected_emergency)
    for row, expected in enumerate(expected_emergency, 2):
        expected_date, expected_period, expected_amount = expected
        actual_date = emergency.cell(row, 1).value
        if expected_date is None:
            emergency_label_errors += int(actual_date is not None)
        else:
            emergency_label_errors += int(_date_text(actual_date) != expected_date)
        emergency_label_errors += int(emergency.cell(row, 2).value != expected_period)
        emergency_amount_diffs.append(float(emergency.cell(row, 3).value or 0) - expected_amount)

    maxima = {
        'max_q_DA_abs_diff_kWh': _max_abs(q_diffs),
        'max_daily_q_DA_total_abs_diff_kWh': _max_abs(q_total_diffs),
        'max_daily_normal_cost_abs_diff_yuan': _max_abs(cost_diffs),
        'max_4h_charge_abs_diff_kWh': _max_abs(charge_diffs),
        'max_4h_discharge_abs_diff_kWh': _max_abs(discharge_diffs),
        'max_SOC_endpoint_abs_diff_kWh': _max_abs(soc_diffs),
        'max_emergency_interval_abs_diff_kWh': _max_abs(emergency_amount_diffs),
    }
    checks.update({
        'purchase_dates': date_errors == 0,
        'purchase_values': max(maxima[k] for k in list(maxima)[:3]) <= TOL,
        'battery_labels': battery_label_errors == 0,
        'battery_values': max(maxima[k] for k in list(maxima)[3:6]) <= TOL,
        'emergency_labels': emergency_label_errors == 0,
        'emergency_values': maxima['max_emergency_interval_abs_diff_kWh'] <= TOL,
    })
    verdict = 'PASS' if all(bool(value) for value in checks.values()) else 'FAIL'
    out = {
        'readback_verdict': verdict, 'checks': checks, **maxima,
        'days': int(df.date.nunique()), 'slot_rows': int(len(df)),
        'emergency_intervals': len(expected_emergency),
        'sheet_dimensions': {name: [wb[name].max_row, wb[name].max_column] for name in SHEETS},
        'source_slot_csv': str(slot_csv), 'workbook': str(workbook),
    }
    output.write_text(json.dumps(out, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(out, ensure_ascii=False, indent=2))
    if verdict != 'PASS':
        raise SystemExit(2)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--slot-csv', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    main(args.slot_csv, args.workbook, args.output)
