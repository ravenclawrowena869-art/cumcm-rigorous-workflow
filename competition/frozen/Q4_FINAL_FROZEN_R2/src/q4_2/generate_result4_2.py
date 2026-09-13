from __future__ import annotations
import argparse, json
from copy import copy
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

PERIODS=['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']

def build(slot_csv: Path, template: Path, output: Path):
    df=pd.read_csv(slot_csv).sort_values(['date','slot'])
    if len(df)!=334*144 or not df.groupby('date').slot.nunique().eq(144).all(): raise RuntimeError('formal causal slot replay coverage failed')
    wb=openpyxl.load_workbook(template)
    s1=wb['计划购电量']; s2=wb['充放电量']; s3=wb['紧急购电量']
    header_style=copy(s1['A1']._style); normal_style=copy(s1['B2']._style)
    s1.delete_rows(2,s1.max_row-1)
    for r,(day,g) in enumerate(df.groupby('date',sort=True),2):
        g=g.sort_values('slot'); s1.cell(r,1,datetime.fromisoformat(day))
        for j,q in enumerate(g.q_DA_kWh,2): s1.cell(r,j,float(q))
        s1.cell(r,146,float(g.q_DA_kWh.sum())); s1.cell(r,147,float(g.normal_cost_yuan.sum()))
        for j in range(1,148): s1.cell(r,j)._style=copy(normal_style)
        s1.cell(r,1).number_format='yyyy-mm-dd'
    s2.delete_rows(2,s2.max_row-1)
    outrow=2
    for day,g in df.groupby('date',sort=True):
        g=g.sort_values('slot')
        for k,label in enumerate(PERIODS):
            b=g.iloc[k*24:(k+1)*24]
            s2.append([datetime.fromisoformat(day) if k==0 else None,label,float(b.charge_exec_kWh.sum()),float(b.discharge_exec_kWh.sum()),'0:00' if k==0 else ('24:00' if k==1 else None),float(g.iloc[0].SOC_start_kWh) if k==0 else (float(g.iloc[-1].SOC_end_kWh) if k==1 else None)])
            for j in range(1,7): s2.cell(outrow,j)._style=copy(normal_style)
            s2.cell(outrow,1).number_format='yyyy-mm-dd'; outrow+=1
    s3.delete_rows(2,s3.max_row-1)
    outrow=2
    for day,g in df.groupby('date',sort=True):
        g=g.sort_values('slot'); active=g[g.emergency_kWh>1e-9].slot.astype(int).tolist(); groups=[]
        for slot in active:
            if not groups or slot!=groups[-1][-1]+1: groups.append([slot])
            else: groups[-1].append(slot)
        first=True
        for x in groups:
            start=10*(x[0]-1); end=10*x[-1]
            fmt=lambda m:'24:00' if m==1440 else f'{m//60}:{m%60:02d}'
            amount=float(g[g.slot.isin(x)].emergency_kWh.sum())
            s3.append([datetime.fromisoformat(day) if first else None,f'{fmt(start)}-{fmt(end)}',amount])
            for j in range(1,4): s3.cell(outrow,j)._style=copy(normal_style)
            s3.cell(outrow,1).number_format='yyyy-mm-dd'; outrow+=1; first=False
    for ws in (s1,s2,s3): ws.freeze_panes='A2'
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)
    return {'status':'PASS','source_slot_csv':str(slot_csv),'output_xlsx':str(output),'days':334,'slot_rows':48096}

if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--slot-csv',type=Path,required=True); ap.add_argument('--template',type=Path,required=True); ap.add_argument('--output',type=Path,required=True)
    a=ap.parse_args(); print(json.dumps(build(a.slot_csv,a.template,a.output),ensure_ascii=False,indent=2))
